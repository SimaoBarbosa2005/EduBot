"""
Document loading and text extraction.

Keeps the old load_all() API, and also exposes load_all_sections() so the
RAG index can preserve source metadata such as PDF page or PPTX slide.
"""

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional


@dataclass(frozen=True)
class DocumentSection:
    source: str
    text: str
    page: Optional[int] = None
    slide: Optional[int] = None


def _try_import_pdf():
    try:
        import pdfplumber

        return pdfplumber
    except ImportError:
        return None


def _try_import_pptx():
    try:
        from pptx import Presentation

        return Presentation
    except ImportError:
        return None


class DocumentLoader:
    """Loads supported files from the documentos/ folder."""

    SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".ppt", ".txt", ".md"}

    def __init__(self, docs_folder: str = "documentos"):
        base = Path(__file__).resolve().parent.parent
        self.docs_folder = base / docs_folder
        self.docs_folder.mkdir(exist_ok=True)

    def list_documents(self) -> list[Path]:
        files = []
        for ext in self.SUPPORTED_EXTENSIONS:
            files.extend(self.docs_folder.glob(f"*{ext}"))
        return sorted(files)

    def load_all(self) -> dict[str, str]:
        docs = {}
        for path in self.list_documents():
            text = self._load_file(path)
            if text:
                docs[path.name] = text
        return docs

    def load_all_sections(self) -> list[DocumentSection]:
        sections = []
        for path in self.list_documents():
            try:
                sections.extend(self._load_file_sections(path))
            except Exception as e:
                print(f"  [Aviso] Erro ao carregar '{path.name}': {e}")
        return sections

    def load_file(self, filename: str) -> Optional[str]:
        path = self.docs_folder / filename
        if not path.exists():
            return None
        return self._load_file(path)

    def _load_file(self, path: Path) -> Optional[str]:
        sections = self._load_file_sections(path)
        if not sections:
            return None

        blocks = []
        for section in sections:
            label = ""
            if section.page is not None:
                label = f"[Pagina {section.page}]\n"
            elif section.slide is not None:
                label = f"[Slide {section.slide}]\n"
            blocks.append(f"{label}{section.text}")
        return "\n\n".join(blocks)

    def _load_file_sections(self, path: Path) -> list[DocumentSection]:
        ext = path.suffix.lower()
        loaders = {
            ".pdf": self._load_pdf_sections,
            ".pptx": self._load_pptx_sections,
            ".ppt": self._load_pptx_sections,
            ".txt": self._load_text_sections,
            ".md": self._load_text_sections,
        }
        loader = loaders.get(ext)
        return loader(path) if loader else []

    def _load_pdf_sections(self, path: Path) -> list[DocumentSection]:
        pdfplumber = _try_import_pdf()
        if pdfplumber is None:
            raise ImportError("pdfplumber nao esta instalado. Execute: pip install pdfplumber")

        sections = []
        with pdfplumber.open(str(path)) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                text = page.extract_text()
                if text:
                    sections.append(DocumentSection(path.name, text.strip(), page=i))
        return sections

    def _load_pptx_sections(self, path: Path) -> list[DocumentSection]:
        Presentation = _try_import_pptx()
        if Presentation is None:
            raise ImportError("python-pptx nao esta instalado. Execute: pip install python-pptx")

        prs = Presentation(str(path))
        sections = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = para.text.strip()
                        if line:
                            texts.append(line)
            if texts:
                sections.append(DocumentSection(path.name, "\n".join(texts), slide=i))
        return sections

    def _load_text_sections(self, path: Path) -> list[DocumentSection]:
        text = self._load_text(path)
        return [DocumentSection(path.name, text)] if text else []

    def _load_text(self, path: Path) -> Optional[str]:
        try:
            return path.read_text(encoding="utf-8").strip() or None
        except UnicodeDecodeError:
            return path.read_text(encoding="latin-1").strip() or None

    @staticmethod
    def clean_text(text: str) -> str:
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r" {2,}", " ", text)
        return text.strip()

    def get_docs_folder_path(self) -> str:
        return str(self.docs_folder)
